"""Extract presidential Electoral College results for 1968 through 2024.

The raw directory mixes National Archives summary workbooks, FEC PDF reports,
modern XLSX workbooks, and legacy XLS workbooks.  This cleaner deliberately
targets presidential Electoral College summaries or allocation tables and
never reads congressional, primary, or national popular-vote summaries as EC
totals.

The 1996 workbook is the one exception to direct EC extraction: it contains
state presidential popular votes only.  Its state winners are combined with
the state EC allocation parsed from the 1992 FEC table (the same post-1990
census apportionment) and checked against the explicit 1996 totals in the 2012
FEC appendices.  The attached 1980 scan contains no usable EC information, so
that single election uses the documented official fallback described below.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
import csv
import hashlib
import math
from numbers import Integral, Real
from pathlib import Path
import re
import shutil
import subprocess
from tempfile import TemporaryDirectory
import unicodedata

import pandas as pd
from pandas.api.types import is_bool_dtype, is_integer_dtype


RAW_DIRECTORY = Path("Raw Data") / "Presidential Election Results"
OUTPUT_RELATIVE_PATH = (
    Path("Cleaned Data")
    / "Presidential Election Results"
    / "Presidential Election Results_cleaned.csv"
)
METADATA_FILENAME = "Access Metadata.csv"
IGNORED_FILENAMES = {".DS_Store", METADATA_FILENAME}

EXPECTED_YEARS = tuple(range(1968, 2025, 4))
TOTAL_ELECTORAL_VOTES = 538
FINAL_COLUMNS = [
    "incumbent_party_ec_votes",
    "challenger_party_ec_votes",
    "incumbent_is_democrat",
]

SOURCE_FILENAMES = {
    1968: "1968 presidential summary.xlsx",
    1972: "1972 presidential summary.xlsx",
    1976: "1976 presidential summary.xlsx",
    1980: "1980 full.pdf",
    1984: "1984 full.pdf",
    1988: "1988 full.pdf",
    1992: "1992 full.pdf",
    1996: "1996 presidential.xlsx",
    2000: "2000 presidential.xls",
    2004: "2004 summary.xls",
    2008: "2008 summary.xls",
    2012: "2012 summary.xls",
    2016: "2016 full.xlsx",
    2020: "2020 full.xlsx",
    2024: "2024 presidential.xlsx",
}

# These are sanity controls, not extraction inputs.  Except for the explicitly
# documented 1980 fallback, every value returned by an extractor must first be
# recovered from a raw source and then agree with this election-specific check.
EXPECTED_PARTY_EC = {
    1968: {"D": 191, "R": 301},
    1972: {"D": 17, "R": 520},
    1976: {"D": 297, "R": 240},
    1980: {"D": 49, "R": 489},
    1984: {"D": 13, "R": 525},
    1988: {"D": 111, "R": 426},
    1992: {"D": 370, "R": 168},
    1996: {"D": 379, "R": 159},
    2000: {"D": 266, "R": 271},
    2004: {"D": 251, "R": 286},
    2008: {"D": 365, "R": 173},
    2012: {"D": 332, "R": 206},
    2016: {"D": 227, "R": 304},
    2020: {"D": 306, "R": 232},
    2024: {"D": 226, "R": 312},
}

INCUMBENT_PARTY = {
    1968: "D",
    1972: "R",
    1976: "R",
    1980: "D",
    1984: "R",
    1988: "R",
    1992: "R",
    1996: "D",
    2000: "D",
    2004: "R",
    2008: "R",
    2012: "D",
    2016: "D",
    2020: "R",
    2024: "D",
}

STATE_NAME_CODES = (
    ("Alabama", "AL"),
    ("Alaska", "AK"),
    ("Arizona", "AZ"),
    ("Arkansas", "AR"),
    ("California", "CA"),
    ("Colorado", "CO"),
    ("Connecticut", "CT"),
    ("Delaware", "DE"),
    ("District of Columbia", "DC"),
    ("Florida", "FL"),
    ("Georgia", "GA"),
    ("Hawaii", "HI"),
    ("Idaho", "ID"),
    ("Illinois", "IL"),
    ("Indiana", "IN"),
    ("Iowa", "IA"),
    ("Kansas", "KS"),
    ("Kentucky", "KY"),
    ("Louisiana", "LA"),
    ("Maine", "ME"),
    ("Maryland", "MD"),
    ("Massachusetts", "MA"),
    ("Michigan", "MI"),
    ("Minnesota", "MN"),
    ("Mississippi", "MS"),
    ("Missouri", "MO"),
    ("Montana", "MT"),
    ("Nebraska", "NE"),
    ("Nevada", "NV"),
    ("New Hampshire", "NH"),
    ("New Jersey", "NJ"),
    ("New Mexico", "NM"),
    ("New York", "NY"),
    ("North Carolina", "NC"),
    ("North Dakota", "ND"),
    ("Ohio", "OH"),
    ("Oklahoma", "OK"),
    ("Oregon", "OR"),
    ("Pennsylvania", "PA"),
    ("Rhode Island", "RI"),
    ("South Carolina", "SC"),
    ("South Dakota", "SD"),
    ("Tennessee", "TN"),
    ("Texas", "TX"),
    ("Utah", "UT"),
    ("Vermont", "VT"),
    ("Virginia", "VA"),
    ("Washington", "WA"),
    ("West Virginia", "WV"),
    ("Wisconsin", "WI"),
    ("Wyoming", "WY"),
)
STATE_CODES = frozenset(code for _, code in STATE_NAME_CODES)

TABLE_SPECS = {
    2000: {
        "sheet": "2000 Pres Elec & Pop Vote-p. 12",
        "aliases": {"D": ("gore",), "R": ("bush",)},
        "title": "2000 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "direct",
        "total_allocation_header": None,
        "state_checks": {"DC": {"D": 2, "R": 0}},
    },
    2004: {
        "sheet": "Table 2. Pres Elec & Pop Vote",
        "aliases": {"D": ("kerry",), "R": ("bush",)},
        "title": "2004 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "direct",
        "total_allocation_header": None,
        "state_checks": {"MN": {"D": 9, "R": 0}},
    },
    2008: {
        "sheet": "Table 2. Electoral &  Pop Vote",
        "aliases": {"D": ("obama",), "R": ("mccain",)},
        "title": "2008 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "stacked",
        "total_allocation_header": None,
        "state_checks": {"NE": {"D": 1, "R": 4}},
    },
    2012: {
        "sheet": "Table 2. Electoral &  Pop Vote",
        "aliases": {"D": ("obama",), "R": ("romney",)},
        "title": "2012 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "stacked",
        "total_allocation_header": None,
        "state_checks": {},
    },
    2016: {
        "sheet": "Table 2. Electoral &  Pop Vote",
        "aliases": {"D": ("clinton",), "R": ("trump",)},
        "title": "2016 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "stacked",
        "total_allocation_header": None,
        "state_checks": {
            "TX": {"D": 0, "R": 36},
            "HI": {"D": 3, "R": 0},
            "WA": {"D": 8, "R": 0},
        },
    },
    2020: {
        "sheet": "3. Table 2 Electoral & Pop Vote",
        # The source misspells both Republican headers as "Trunp (R)".
        "aliases": {"D": ("biden",), "R": ("trump", "trunp")},
        "title": "2020 presidential electoral and popular vote",
        "max_rows": 70,
        "max_cols": 10,
        "header_style": "stacked",
        "total_allocation_header": None,
        "state_checks": {
            "ME": {"D": 3, "R": 1},
            "NE": {"D": 1, "R": 4},
        },
    },
    2024: {
        "sheet": "OFFICIAL 2024 PRES GE RESULTS",
        "aliases": {"D": ("harris",), "R": ("trump",)},
        "title": None,
        "max_rows": 70,
        "max_cols": 35,
        "header_style": "direct",
        "total_allocation_header": "electoral votes",
        "state_checks": {
            "ME": {"D": 3, "R": 1},
            "NE": {"D": 1, "R": 4},
        },
    },
}

PDF_PAGE_COUNTS = {1980: 5, 1984: 123, 1988: 80, 1992: 126}
PDF_TARGET_PAGES = {1984: 19, 1988: 21, 1992: 14}
PDF_1992_CONFIRMATION_PAGE = 15
EXPECTED_1980_SHA256 = (
    "b77a61c54d19e68f11a2b3f8272ec08f2141d7faaed5e57b4453bf7c0bac35ad"
)


class ExtractionError(ValueError):
    """Raised when a source does not support one unambiguous EC extraction."""


def normalize_text(value: object) -> str:
    """Normalize source text while retaining punctuation useful to parsers."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    text = text.translate(str.maketrans({"–": "-", "—": "-", "−": "-"}))
    return re.sub(r"\s+", " ", text).strip().casefold()


def row_text(row: Sequence[object]) -> str:
    """Return one normalized string for a possibly merged/sparse source row."""
    return normalize_text(
        " ".join(str(value) for value in row if normalize_text(value))
    )


def first_nonblank(row: Sequence[object]) -> object | None:
    """Return the first source cell that is not blank after normalization."""
    return next((value for value in row if normalize_text(value)), None)


def parse_strict_int(value: object, context: str) -> int:
    """Parse an integer, allowing commas, brackets, and trailing footnote marks."""
    if isinstance(value, bool):
        raise ExtractionError(f"{context}: boolean is not an Electoral College count.")
    if isinstance(value, Integral):
        return int(value)
    if isinstance(value, Real):
        number = float(value)
        if math.isfinite(number) and number.is_integer():
            return int(number)
        raise ExtractionError(f"{context}: expected an integer; found {value!r}.")

    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ").strip()
    match = re.fullmatch(
        r"[\[(]?\s*([0-9]+(?:,[0-9]{3})*)\s*[\])]?\s*[\*#†‡]*",
        text,
    )
    if match is None:
        raise ExtractionError(f"{context}: cannot normalize integer value {value!r}.")
    return int(match.group(1).replace(",", ""))


def optional_strict_int(value: object) -> int | None:
    """Return an integer when a cell is strictly parseable, otherwise None."""
    if not normalize_text(value):
        return None
    try:
        return parse_strict_int(value, "candidate total")
    except ExtractionError:
        return None


def select_sheet_name(available: Sequence[str], expected: str, path: Path) -> str:
    """Select one sheet, tolerating whitespace/case differences but not ambiguity."""
    if expected in available:
        return expected
    normalized_expected = normalize_text(expected)
    matches = [
        name for name in available if normalize_text(name) == normalized_expected
    ]
    if len(matches) != 1:
        raise ExtractionError(
            f"{path.name}: expected one sheet matching {expected!r}; "
            f"found {matches or 'none'}. Available sheets: {list(available)}"
        )
    return matches[0]


def trim_rows(rows: Iterable[Sequence[object]]) -> list[list[object]]:
    """Drop trailing blank cells but retain row positions and internal blanks."""
    result: list[list[object]] = []
    for source_row in rows:
        row = list(source_row)
        while row and not normalize_text(row[-1]):
            row.pop()
        result.append(row)
    return result


def read_openpyxl_sheet(
    path: Path, sheet_name: str, max_rows: int, max_cols: int
) -> list[list[object]]:
    """Read a bounded XLSX range without traversing unrelated inflated sheets."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Reading .xlsx files requires the openpyxl package."
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        selected = select_sheet_name(workbook.sheetnames, sheet_name, path)
        worksheet = workbook[selected]
        rows = worksheet.iter_rows(
            min_row=1,
            max_row=max_rows,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
        return trim_rows(rows)
    finally:
        workbook.close()


def read_xlrd_sheet(
    path: Path, sheet_name: str, max_rows: int, max_cols: int
) -> list[list[object]]:
    """Read a bounded legacy XLS sheet with xlrd when it is installed."""
    import xlrd  # type: ignore[import-not-found]

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        selected = select_sheet_name(workbook.sheet_names(), sheet_name, path)
        worksheet = workbook.sheet_by_name(selected)
        rows = (
            [
                worksheet.cell_value(row_index, column_index)
                for column_index in range(min(worksheet.ncols, max_cols))
            ]
            for row_index in range(min(worksheet.nrows, max_rows))
        )
        return trim_rows(rows)
    finally:
        workbook.release_resources()


def read_xls_via_libreoffice(
    path: Path, sheet_name: str, max_rows: int, max_cols: int
) -> list[list[object]]:
    """Convert one XLS to a temporary XLSX if a direct XLS engine is unavailable."""
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        raise RuntimeError(
            f"Cannot read legacy workbook {path.name}: install xlrd or make "
            "LibreOffice (soffice) available on PATH."
        )

    with TemporaryDirectory(prefix="presidential-ec-xls-") as output_directory:
        with TemporaryDirectory(prefix="presidential-ec-profile-") as profile_directory:
            command = [
                executable,
                f"-env:UserInstallation={Path(profile_directory).as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                output_directory,
                str(path),
            ]
            try:
                completed = subprocess.run(
                    command,
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=90,
                )
            except subprocess.TimeoutExpired as error:
                raise RuntimeError(
                    f"LibreOffice timed out while reading {path.name}."
                ) from error

            converted_files = list(Path(output_directory).glob("*.xlsx"))
            if completed.returncode != 0 or len(converted_files) != 1:
                details = (completed.stderr or completed.stdout).strip()
                raise RuntimeError(
                    f"LibreOffice could not convert {path.name} to temporary XLSX. "
                    f"Return code {completed.returncode}; output: {details!r}"
                )
            return read_openpyxl_sheet(
                converted_files[0], sheet_name, max_rows, max_cols
            )


def read_spreadsheet_sheet(
    path: Path, sheet_name: str, max_rows: int, max_cols: int
) -> list[list[object]]:
    """Read a bounded sheet from XLSX or legacy XLS with a clear fallback."""
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        return read_openpyxl_sheet(path, sheet_name, max_rows, max_cols)
    if suffix != ".xls":
        raise ExtractionError(f"Unsupported spreadsheet format: {path.name}")

    try:
        return read_xlrd_sheet(path, sheet_name, max_rows, max_cols)
    except (ImportError, ModuleNotFoundError):
        return read_xls_via_libreoffice(path, sheet_name, max_rows, max_cols)
    except Exception as direct_error:
        try:
            return read_xls_via_libreoffice(path, sheet_name, max_rows, max_cols)
        except Exception as conversion_error:
            raise RuntimeError(
                f"Both direct XLS reading and temporary LibreOffice conversion "
                f"failed for {path.name}. Direct error: {direct_error}; "
                f"conversion error: {conversion_error}"
            ) from conversion_error


def discover_sources(repo_root: Path) -> tuple[Path, dict[int, Path]]:
    """Require exactly the documented election source file for every year."""
    raw_directory = repo_root / RAW_DIRECTORY
    if not raw_directory.is_dir():
        raise FileNotFoundError(f"Raw data directory not found: {raw_directory}")

    expected_names = set(SOURCE_FILENAMES.values())
    actual_names = {
        path.name
        for path in raw_directory.iterdir()
        if path.is_file() and path.name not in IGNORED_FILENAMES
    }
    missing = sorted(expected_names - actual_names)
    unexpected = sorted(actual_names - expected_names)
    if missing or unexpected:
        raise ExtractionError(
            "Election source inventory differs from the expected 1968-2024 "
            f"corpus. Missing: {missing or 'none'}; "
            f"unexpected: {unexpected or 'none'}."
        )

    sources = {year: raw_directory / name for year, name in SOURCE_FILENAMES.items()}
    if tuple(sorted(sources)) != EXPECTED_YEARS:
        raise AssertionError("Internal source-year mapping is incomplete.")
    validate_access_metadata(raw_directory / METADATA_FILENAME, expected_names)
    return raw_directory, sources


def validate_access_metadata(metadata_path: Path, expected_names: set[str]) -> None:
    """Verify provenance metadata covers each election source exactly once."""
    if not metadata_path.is_file():
        raise FileNotFoundError(f"Access metadata not found: {metadata_path}")
    with metadata_path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        expected_fields = ["file_name", "url", "access_date"]
        if reader.fieldnames != expected_fields:
            raise ExtractionError(
                f"{metadata_path.name}: expected columns {expected_fields}; "
                f"found {reader.fieldnames}."
            )
        records = list(reader)

    names = [record["file_name"].strip() for record in records]
    if len(names) != len(set(names)):
        raise ExtractionError(f"{metadata_path.name} contains duplicate file names.")
    if set(names) != expected_names:
        raise ExtractionError(
            f"{metadata_path.name} does not cover the exact election file inventory."
        )
    for record in records:
        if not record["url"].strip():
            raise ExtractionError(
                f"{metadata_path.name}: missing URL for {record['file_name']}."
            )
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", record["access_date"].strip()) is None:
            raise ExtractionError(
                f"{metadata_path.name}: invalid access date for {record['file_name']}."
            )


def unique_labeled_value(
    rows: Sequence[Sequence[object]], label_predicate, context: str
) -> object:
    """Return the sole nonblank value following a uniquely matched row label."""
    matches: list[tuple[int, Sequence[object]]] = []
    for row_number, row in enumerate(rows, start=1):
        label = normalize_text(first_nonblank(row))
        if label and label_predicate(label):
            matches.append((row_number, row))
    if len(matches) != 1:
        raise ExtractionError(
            f"{context}: expected exactly one matching label row; found {len(matches)}."
        )
    row_number, row = matches[0]
    nonblank = [value for value in row if normalize_text(value)]
    if len(nonblank) != 2:
        raise ExtractionError(
            f"{context}, row {row_number}: expected one label and one value; "
            f"found {nonblank}."
        )
    return nonblank[1]


def candidate_party(candidate: object, context: str) -> str:
    """Extract a Democratic/Republican marker from a candidate label."""
    match = re.search(r"\[\s*([DR])\s*\]", str(candidate), flags=re.IGNORECASE)
    if match is None:
        raise ExtractionError(f"{context}: candidate label has no [D]/[R] marker.")
    return match.group(1).upper()


def extract_nara_summary(year: int, path: Path) -> dict[str, int]:
    """Extract the explicit national winner/opponent totals from a NARA sheet."""
    rows = read_spreadsheet_sheet(path, "Summary", max_rows=30, max_cols=5)
    if not rows or f"{year} electoral college results" not in row_text(rows[0]):
        raise ExtractionError(f"{path.name}: unexpected or missing election title.")

    winner = unique_labeled_value(
        rows, lambda label: label == "president", f"{path.name} winner"
    )
    opponent = unique_labeled_value(
        rows, lambda label: label == "main opponent", f"{path.name} opponent"
    )
    winner_votes = parse_strict_int(
        unique_labeled_value(
            rows,
            lambda label: "electoral vote" in label and "winner" in label,
            f"{path.name} winner EC",
        ),
        f"{path.name} winner EC",
    )
    opponent_votes = parse_strict_int(
        unique_labeled_value(
            rows,
            lambda label: "electoral vote" in label and "main opponent" in label,
            f"{path.name} opponent EC",
        ),
        f"{path.name} opponent EC",
    )
    source_total = parse_strict_int(
        unique_labeled_value(
            rows,
            lambda label: "electoral vote" in label and "total" in label,
            f"{path.name} total EC",
        ),
        f"{path.name} total EC",
    )
    if source_total != TOTAL_ELECTORAL_VOTES:
        raise ExtractionError(
            f"{path.name}: explicit EC total is {source_total}, not 538."
        )

    result = {
        candidate_party(winner, f"{path.name} winner"): winner_votes,
        candidate_party(opponent, f"{path.name} opponent"): opponent_votes,
    }
    if set(result) != {"D", "R"}:
        raise ExtractionError(
            f"{path.name}: major-party labels are ambiguous: {result}."
        )

    other_value = unique_labeled_value(
        rows,
        lambda label: label in {"other presidential vote", "votes for others"},
        f"{path.name} other presidential vote",
    )
    other_matches = re.findall(r"\(\s*(\d+)\s*\)", str(other_value))
    if len(other_matches) != 1:
        raise ExtractionError(
            f"{path.name}: cannot identify one other presidential EC total in "
            f"{other_value!r}."
        )
    source_other = int(other_matches[0])
    residual = source_total - result["D"] - result["R"]
    if source_other != residual:
        raise ExtractionError(
            f"{path.name}: explicit other EC vote {source_other} disagrees with "
            f"the residual {residual}."
        )
    return result


def load_pdf(path: Path, year: int):
    """Open a PDF and enforce the inspected source's physical page count."""
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError(
            "Reading election PDFs requires the pypdf package."
        ) from error

    reader = PdfReader(path)
    expected_pages = PDF_PAGE_COUNTS[year]
    if len(reader.pages) != expected_pages:
        raise ExtractionError(
            f"{path.name}: expected {expected_pages} pages; found {len(reader.pages)}. "
            "Review the source before changing a targeted-page extractor."
        )
    return reader


def one_regex_match(pattern: str, text: str, context: str) -> tuple[str, ...]:
    """Require exactly one regex match and return its capture groups."""
    matches = re.findall(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    if len(matches) != 1:
        raise ExtractionError(
            f"{context}: expected exactly one summary match; found {len(matches)}."
        )
    match = matches[0]
    return (match,) if isinstance(match, str) else tuple(match)


def extract_1980_fallback(path: Path) -> dict[str, int]:
    """Return the sole documented fallback after fingerprinting the deficient PDF."""
    reader = load_pdf(path, 1980)
    del reader
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != EXPECTED_1980_SHA256:
        raise ExtractionError(
            f"{path.name}: source fingerprint changed. The 1980 fallback is "
            "authorized only for the inspected five-page FEC popular-vote scan."
        )

    # The attached 1980 FEC scan contains national and state presidential
    # popular-vote totals, but no usable national or state Electoral College
    # allocation table. Use the official cast-vote result: Reagan 489, Carter 49.
    # Official result: https://www.archives.gov/electoral-college/1980
    return {"R": 489, "D": 49}


def extract_1984_pdf(path: Path) -> dict[str, int]:
    """Extract the candidate-labelled EC line on physical page 20."""
    reader = load_pdf(path, 1984)
    text = reader.pages[PDF_TARGET_PAGES[1984]].extract_text() or ""
    reagan, mondale = one_regex_match(
        r"Electoral\s+Vote\s+Reagan\s*-\s*(\d+)\s+Mondale\s*-\s*(\d+)",
        text,
        f"{path.name}, physical page 20",
    )
    return {"R": int(reagan), "D": int(mondale)}


def ocr_integer(token: str, context: str) -> int:
    """Normalize the inspected OCR substitutions used in the 1988 summary."""
    normalized = token.translate(str.maketrans({"l": "1", "I": "1", "|": "1"}))
    if re.fullmatch(r"\d+", normalized) is None:
        raise ExtractionError(f"{context}: invalid OCR integer token {token!r}.")
    return int(normalized)


def extract_1988_pdf(path: Path) -> dict[str, int]:
    """Extract the OCR-backed EC line on physical page 22."""
    reader = load_pdf(path, 1988)
    text = reader.pages[PDF_TARGET_PAGES[1988]].extract_text() or ""
    bush_token, dukakis_token, bentsen_token = one_regex_match(
        r"Bush\s*-\s*([0-9Il|]+).{0,30}?\S*kakis\s*-\s*"
        r"([0-9Il|]+)\s+Bentsen\s*-\s*([0-9Il|]+)",
        text,
        f"{path.name}, physical page 22",
    )
    bush = ocr_integer(bush_token, f"{path.name} Bush total")
    dukakis = ocr_integer(dukakis_token, f"{path.name} Dukakis total")
    bentsen = ocr_integer(bentsen_token, f"{path.name} Bentsen total")
    if bentsen != TOTAL_ELECTORAL_VOTES - bush - dukakis:
        raise ExtractionError(
            f"{path.name}: Bentsen EC vote does not equal the national residual."
        )
    return {"R": bush, "D": dukakis}


def extract_1992_pdf(path: Path) -> tuple[dict[str, int], dict[str, int]]:
    """Extract national totals and the state allocation on physical page 15."""
    reader = load_pdf(path, 1992)
    summary_text = reader.pages[PDF_TARGET_PAGES[1992]].extract_text() or ""
    if "1992 electoral and popular vote summary" not in normalize_text(summary_text):
        raise ExtractionError(f"{path.name}: targeted page lacks the 1992 EC title.")
    if re.search(r"Bush\s+Clinton", summary_text, flags=re.IGNORECASE) is None:
        raise ExtractionError(f"{path.name}: EC candidate header order is ambiguous.")

    bush, clinton = one_regex_match(
        r"Total:\s*(\d+)\s+(\d+)",
        summary_text,
        f"{path.name}, physical page 15 total row",
    )
    result = {"R": int(bush), "D": int(clinton)}

    allocation_matches = re.findall(
        r"^\s*([A-Za-z]{2})\s+(\d{1,2})\s+",
        summary_text,
        flags=re.MULTILINE,
    )
    allocations = {state.upper(): int(votes) for state, votes in allocation_matches}
    if len(allocation_matches) != len(allocations):
        raise ExtractionError(f"{path.name}: duplicate state allocation rows found.")
    if set(allocations) != STATE_CODES or sum(allocations.values()) != 538:
        raise ExtractionError(
            f"{path.name}: state EC allocation is incomplete or does not total 538."
        )

    confirmation = reader.pages[PDF_1992_CONFIRMATION_PAGE].extract_text() or ""
    confirmed_clinton = int(
        one_regex_match(
            r"Clinton\s*\(Democrat\)\s*(\d+)",
            confirmation,
            f"{path.name}, physical page 16 Clinton confirmation",
        )[0]
    )
    confirmed_bush = int(
        one_regex_match(
            r"Bush\s*\(Republican\)\s*(\d+)",
            confirmation,
            f"{path.name}, physical page 16 Bush confirmation",
        )[0]
    )
    if result != {"D": confirmed_clinton, "R": confirmed_bush}:
        raise ExtractionError(
            f"{path.name}: summary total row disagrees with the candidate-labelled "
            "distribution on the next page."
        )
    return result, allocations


def candidate_popular_vote(
    segment: Sequence[Sequence[object]], surname: str, context: str
) -> int:
    """Recover a major candidate's state vote despite irregular merged cells."""
    candidates: list[int] = []
    for row in segment:
        if re.search(rf"\b{re.escape(surname)}\b", row_text(row)) is None:
            continue
        for value in row:
            if isinstance(value, bool):
                continue
            if isinstance(value, Integral):
                candidates.append(int(value))
            elif isinstance(value, Real):
                number = float(value)
                if math.isfinite(number) and number.is_integer():
                    candidates.append(int(number))
            elif isinstance(value, str):
                candidates.extend(
                    int(match.replace(",", ""))
                    for match in re.findall(
                        r"(?<![\d.])\d{1,3}(?:,\d{3})+(?![\d.])", value
                    )
                )
    candidates = [value for value in candidates if value > TOTAL_ELECTORAL_VOTES]
    if not candidates:
        raise ExtractionError(f"{context}: no unambiguous popular-vote count found.")
    return max(candidates)


def appendix_candidate_total(path: Path, sheet_name: str, header: str) -> int:
    """Read one candidate-year EC total from an FEC ranking appendix."""
    rows = read_spreadsheet_sheet(path, sheet_name, max_rows=75, max_cols=30)
    positions = [
        (row_index, column_index)
        for row_index, row in enumerate(rows)
        for column_index, value in enumerate(row)
        if normalize_text(value) == normalize_text(header)
    ]
    if len(positions) != 1:
        raise ExtractionError(
            f"{path.name}/{sheet_name}: expected one {header!r} header; "
            f"found {len(positions)}."
        )
    header_row, column = positions[0]
    total_rows = [
        row_index
        for row_index in range(header_row + 1, len(rows))
        if "total electoral votes" in row_text(rows[row_index])
    ]
    if len(total_rows) != 1:
        raise ExtractionError(
            f"{path.name}/{sheet_name}: expected one total EC row; "
            f"found {len(total_rows)}."
        )
    total_row = rows[total_rows[0]]
    if column >= len(total_row):
        raise ExtractionError(
            f"{path.name}/{sheet_name}: total cell for {header!r} is blank."
        )
    return parse_strict_int(
        total_row[column], f"{path.name}/{sheet_name} {header} total"
    )


def extract_1996_state_results(
    path: Path, allocations: dict[str, int], appendix_path: Path
) -> dict[str, int]:
    """Derive 1996 EC totals from state winners and verify FEC appendix totals."""
    rows = read_spreadsheet_sheet(path, "Table 1", max_rows=600, max_cols=18)
    if not rows or "1996 presidential general election results" not in row_text(
        rows[0]
    ):
        raise ExtractionError(f"{path.name}: unexpected or missing election title.")

    total_row_indices = []
    for row_index, row in enumerate(rows):
        first = normalize_text(first_nonblank(row))
        if re.fullmatch(r"total(?:\s+votes?)?(?::.*)?", first):
            total_row_indices.append(row_index)
    if len(total_row_indices) != len(STATE_NAME_CODES):
        raise ExtractionError(
            f"{path.name}: expected 51 state/DC result blocks; "
            f"found {len(total_row_indices)}."
        )

    totals = {"D": 0, "R": 0}
    segment_start = 0
    for (state_name, state_code), segment_end in zip(
        STATE_NAME_CODES, total_row_indices, strict=True
    ):
        segment = rows[segment_start : segment_end + 1]
        segment_start = segment_end + 1
        if normalize_text(state_name) not in row_text(
            [value for row in segment for value in row]
        ):
            raise ExtractionError(
                f"{path.name}: state block ending on row {segment_end + 1} "
                f"does not identify expected state {state_name}."
            )
        clinton = candidate_popular_vote(
            segment, "clinton", f"{path.name} {state_name} Clinton"
        )
        dole = candidate_popular_vote(segment, "dole", f"{path.name} {state_name} Dole")
        if clinton == dole:
            raise ExtractionError(
                f"{path.name}: tied major-party vote in {state_name}."
            )
        winner_party = "D" if clinton > dole else "R"
        totals[winner_party] += allocations[state_code]

    # Maine and Nebraska can split electors by congressional district.  The
    # state-winner reconstruction is accepted only because the independent FEC
    # appendix totals below agree exactly, confirming no 1996 split changes it.
    appendix_totals = {
        "D": appendix_candidate_total(appendix_path, "Appendix B", "CLINTON 1996"),
        "R": appendix_candidate_total(appendix_path, "Appendix C", "DOLE 1996"),
    }
    if totals != appendix_totals:
        raise ExtractionError(
            f"{path.name}: state-winner reconstruction {totals} disagrees with "
            f"the explicit 2012 FEC appendix totals {appendix_totals}."
        )
    return totals


def state_code_from_cell(value: object) -> str | None:
    """Normalize state abbreviations carrying source footnote markers."""
    match = re.fullmatch(r"\s*([A-Za-z]{2})\s*[\*#†‡]*\s*", str(value or ""))
    if match is None:
        return None
    code = match.group(1).upper()
    return code if code in STATE_CODES else None


def validate_ec_header_semantics(
    path: Path,
    rows: Sequence[Sequence[object]],
    first_state_index: int,
    total_row: Sequence[object],
    columns: dict[str, int],
    header_positions: dict[str, list[tuple[int, int, object]]],
    header_style: str,
) -> None:
    """Prove selected candidate columns belong to the Electoral Vote block."""
    selected_positions: dict[str, list[tuple[int, int, object]]] = {
        party: [
            position
            for position in header_positions[party]
            if position[1] == columns[party]
        ]
        for party in ("D", "R")
    }
    if any(len(positions) != 1 for positions in selected_positions.values()):
        raise ExtractionError(
            f"{path.name}: each selected EC column must have exactly one "
            "candidate/party header."
        )

    if header_style == "direct":
        for party in ("D", "R"):
            header = normalize_text(selected_positions[party][0][2])
            if "electoral vote" not in header or "popular vote" in header:
                raise ExtractionError(
                    f"{path.name}: selected {party} column is not explicitly "
                    "labelled as an Electoral Vote column."
                )
        return

    if header_style != "stacked":
        raise AssertionError(f"Unsupported table header style: {header_style!r}")

    selected_columns = sorted(columns.values())
    if selected_columns[1] - selected_columns[0] != 1:
        raise ExtractionError(
            f"{path.name}: stacked major-party EC columns are not adjacent."
        )
    selected_rows = {positions[0][0] for positions in selected_positions.values()}
    if len(selected_rows) != 1:
        raise ExtractionError(
            f"{path.name}: stacked major-party candidate headers are misaligned."
        )
    candidate_row = selected_rows.pop()
    for positions in selected_positions.values():
        if "popular vote" in normalize_text(positions[0][2]):
            raise ExtractionError(
                f"{path.name}: selected EC candidate header says Popular Vote."
            )

    parent_rows = rows[:candidate_row]
    electoral_anchor_rows = [
        row_index
        for row_index, row in enumerate(parent_rows)
        if selected_columns[0] < len(row)
        and "electoral vote" in normalize_text(row[selected_columns[0]])
        and "popular vote" not in normalize_text(row[selected_columns[0]])
    ]
    if not electoral_anchor_rows:
        raise ExtractionError(
            f"{path.name}: no Electoral Vote group header anchors the selected "
            "two-column candidate block."
        )
    # Use the nearest structural group-header row, not a title farther above the
    # table that happens to contain both phrases.
    group_header_row = parent_rows[max(electoral_anchor_rows)]

    # Both parties also appear in the later popular-vote block.  Requiring that
    # parallel block makes the left-hand EC selection semantic rather than a
    # coincidence based only on the candidate total being less than 538.
    popular_candidate_columns: list[int] = []
    for party in ("D", "R"):
        popular_positions = []
        for _, column_index, _ in header_positions[party]:
            if column_index <= selected_columns[1] or column_index >= len(total_row):
                continue
            total_value = optional_strict_int(total_row[column_index])
            if total_value is not None and total_value > TOTAL_ELECTORAL_VOTES:
                popular_positions.append(column_index)
        if len(set(popular_positions)) != 1:
            raise ExtractionError(
                f"{path.name}: cannot distinguish one later {party} Popular Vote "
                "column from the selected EC block."
            )
        popular_candidate_columns.extend(popular_positions)

    popular_block_start = min(popular_candidate_columns)
    if not any(
        selected_columns[1] <= column_index <= popular_block_start
        and "popular vote" in normalize_text(value)
        for column_index, value in enumerate(group_header_row)
    ):
        raise ExtractionError(
            f"{path.name}: the Electoral Vote structural header row has no "
            "matching Popular Vote group header."
        )


def extract_ec_table(year: int, path: Path) -> dict[str, int]:
    """Extract one bounded presidential EC allocation table and its total row."""
    spec = TABLE_SPECS[year]
    rows = read_spreadsheet_sheet(
        path,
        str(spec["sheet"]),
        max_rows=int(spec["max_rows"]),
        max_cols=int(spec["max_cols"]),
    )
    if not rows:
        raise ExtractionError(f"{path.name}: presidential EC sheet is empty.")
    expected_title = spec["title"]
    if expected_title is not None and normalize_text(expected_title) not in row_text(
        rows[0]
    ):
        raise ExtractionError(f"{path.name}: unexpected presidential EC table title.")

    total_rows = [
        index
        for index, row in enumerate(rows)
        if re.fullmatch(r"total\s*:?", normalize_text(first_nonblank(row)))
    ]
    if len(total_rows) != 1:
        raise ExtractionError(
            f"{path.name}: expected one national Total row in the presidential "
            f"EC sheet; found {len(total_rows)}."
        )
    total_index = total_rows[0]
    total_row = rows[total_index]

    state_rows: list[tuple[int, str]] = []
    for row_index, row in enumerate(rows[:total_index]):
        if not row:
            continue
        state_code = state_code_from_cell(row[0])
        if state_code is not None:
            state_rows.append((row_index, state_code))
    state_code_list = [state for _, state in state_rows]
    if len(state_code_list) != 51 or set(state_code_list) != STATE_CODES:
        raise ExtractionError(
            f"{path.name}: presidential EC table must contain each state/DC once; "
            f"found {len(state_code_list)} recognized rows."
        )
    if len(state_code_list) != len(set(state_code_list)):
        raise ExtractionError(f"{path.name}: duplicate state rows in EC table.")
    first_state_index = state_rows[0][0]

    columns: dict[str, int] = {}
    header_positions: dict[str, list[tuple[int, int, object]]] = {}
    aliases_by_party = spec["aliases"]
    for party in ("D", "R"):
        aliases = tuple(aliases_by_party[party])
        party_pattern = re.compile(rf"\(\s*{party}\s*\)", flags=re.IGNORECASE)
        positions = [
            (row_index, column_index, value)
            for row_index, row in enumerate(rows[:first_state_index])
            for column_index, value in enumerate(row)
            if party_pattern.search(str(value))
            and any(alias in normalize_text(value) for alias in aliases)
        ]
        header_positions[party] = positions
        header_columns = {column_index for _, column_index, _ in positions}
        ec_columns = []
        for column_index in header_columns:
            if column_index >= len(total_row):
                continue
            total_value = optional_strict_int(total_row[column_index])
            if total_value is not None and 0 <= total_value <= TOTAL_ELECTORAL_VOTES:
                ec_columns.append(column_index)
        if len(ec_columns) != 1:
            raise ExtractionError(
                f"{path.name}: expected one {party} candidate EC column after "
                f"excluding popular-vote columns; found {ec_columns}."
            )
        columns[party] = ec_columns[0]

    if columns["D"] == columns["R"]:
        raise ExtractionError(f"{path.name}: major-party EC columns overlap.")
    validate_ec_header_semantics(
        path,
        rows,
        first_state_index,
        total_row,
        columns,
        header_positions,
        str(spec["header_style"]),
    )

    explicit_totals = {
        party: parse_strict_int(
            total_row[column], f"{path.name} {party} national EC total"
        )
        for party, column in columns.items()
    }
    state_allocations: dict[str, dict[str, int]] = {}
    for row_index, state_code in state_rows:
        row = rows[row_index]
        state_allocations[state_code] = {}
        for party, column in columns.items():
            value = row[column] if column < len(row) else None
            if normalize_text(value):
                parsed = parse_strict_int(
                    value, f"{path.name} {state_code} {party} EC allocation"
                )
            else:
                parsed = 0
            if not 0 <= parsed <= 55:
                raise ExtractionError(
                    f"{path.name}: implausible {state_code} {party} EC value {parsed}."
                )
            state_allocations[state_code][party] = parsed

    summed_totals = {
        party: sum(allocation[party] for allocation in state_allocations.values())
        for party in ("D", "R")
    }
    if summed_totals != explicit_totals:
        raise ExtractionError(
            f"{path.name}: state EC sums {summed_totals} disagree with the "
            f"explicit national row {explicit_totals}."
        )

    total_allocation_header = spec["total_allocation_header"]
    if total_allocation_header is not None:
        total_header_positions = [
            (row_index, column_index)
            for row_index, row in enumerate(rows[:first_state_index])
            for column_index, value in enumerate(row)
            if normalize_text(value) == normalize_text(total_allocation_header)
        ]
        if len(total_header_positions) != 1:
            raise ExtractionError(
                f"{path.name}: expected one {total_allocation_header!r} total "
                f"allocation header; found {len(total_header_positions)}."
            )
        _, allocation_column = total_header_positions[0]
        if allocation_column >= len(total_row):
            raise ExtractionError(
                f"{path.name}: national Electoral Votes total cell is blank."
            )
        explicit_allocation_total = parse_strict_int(
            total_row[allocation_column],
            f"{path.name} explicit national Electoral Votes total",
        )
        possible_by_state: dict[str, int] = {}
        for row_index, state_code in state_rows:
            row = rows[row_index]
            if allocation_column >= len(row) or not normalize_text(
                row[allocation_column]
            ):
                raise ExtractionError(
                    f"{path.name}: missing total EC allocation for {state_code}."
                )
            possible = parse_strict_int(
                row[allocation_column],
                f"{path.name} {state_code} total EC allocation",
            )
            if not 1 <= possible <= 55:
                raise ExtractionError(
                    f"{path.name}: implausible total EC allocation for "
                    f"{state_code}: {possible}."
                )
            possible_by_state[state_code] = possible
            major_allocated = sum(state_allocations[state_code].values())
            if major_allocated > possible:
                raise ExtractionError(
                    f"{path.name}: major-party EC allocations exceed the "
                    f"{state_code} total."
                )
        if explicit_allocation_total != TOTAL_ELECTORAL_VOTES:
            raise ExtractionError(
                f"{path.name}: explicit national Electoral Votes total is "
                f"{explicit_allocation_total}, not 538."
            )
        if sum(possible_by_state.values()) != explicit_allocation_total:
            raise ExtractionError(
                f"{path.name}: state total EC allocations do not reconcile "
                "to the explicit national total."
            )

    source_total_lines = [
        row_text(row)
        for row in rows[total_index + 1 :]
        if "total electoral vote" in row_text(row)
        and re.search(r"\b538\b", row_text(row))
    ]
    if len(source_total_lines) != 1:
        raise ExtractionError(
            f"{path.name}: expected one source statement confirming 538 total "
            f"electors; found {len(source_total_lines)}."
        )

    state_checks = spec["state_checks"]
    for state_code, expected in state_checks.items():
        if state_allocations[state_code] != expected:
            raise ExtractionError(
                f"{path.name}: {state_code} EC allocation "
                f"{state_allocations[state_code]} disagrees with expected {expected}."
            )
    return explicit_totals


def validate_source_result(year: int, result: dict[str, int]) -> None:
    """Apply the election-specific national sanity control after extraction."""
    if set(result) != {"D", "R"}:
        raise ExtractionError(f"{year}: extracted party keys are {sorted(result)}.")
    if any(
        isinstance(value, bool) or not isinstance(value, Integral)
        for value in result.values()
    ):
        raise ExtractionError(f"{year}: EC values must be integers; found {result}.")
    result = {party: int(value) for party, value in result.items()}
    if result != EXPECTED_PARTY_EC[year]:
        raise ExtractionError(
            f"{year}: extracted party totals {result} disagree with the explicit "
            f"election sanity control {EXPECTED_PARTY_EC[year]}."
        )
    residual = TOTAL_ELECTORAL_VOTES - result["D"] - result["R"]
    if residual < 0:
        raise ExtractionError(f"{year}: major-party EC totals exceed 538.")


def extract_all_results(sources: dict[int, Path]) -> dict[int, dict[str, int]]:
    """Run each source-appropriate extractor and retain party-level EC totals."""
    results: dict[int, dict[str, int]] = {}
    for year in (1968, 1972, 1976):
        results[year] = extract_nara_summary(year, sources[year])

    results[1980] = extract_1980_fallback(sources[1980])
    results[1984] = extract_1984_pdf(sources[1984])
    results[1988] = extract_1988_pdf(sources[1988])
    results[1992], allocations = extract_1992_pdf(sources[1992])
    results[1996] = extract_1996_state_results(
        sources[1996], allocations, sources[2012]
    )

    for year in (2000, 2004, 2008, 2012, 2016, 2020, 2024):
        results[year] = extract_ec_table(year, sources[year])

    if tuple(sorted(results)) != EXPECTED_YEARS:
        raise AssertionError("Not every expected election was extracted.")
    for year in EXPECTED_YEARS:
        validate_source_result(year, results[year])
    return results


def build_cleaned_frame(
    results: dict[int, dict[str, int]],
) -> tuple[pd.DataFrame, pd.Series]:
    """Map party totals to incumbent/challenger semantics and validate output."""
    records = []
    other_values: dict[int, int] = {}
    for year in EXPECTED_YEARS:
        incumbent_party = INCUMBENT_PARTY[year]
        challenger_party = "R" if incumbent_party == "D" else "D"
        incumbent_votes = int(results[year][incumbent_party])
        challenger_votes = int(results[year][challenger_party])
        other_votes = TOTAL_ELECTORAL_VOTES - incumbent_votes - challenger_votes
        records.append(
            {
                "year": year,
                "incumbent_party_ec_votes": incumbent_votes,
                "challenger_party_ec_votes": challenger_votes,
                "incumbent_is_democrat": incumbent_party == "D",
            }
        )
        other_values[year] = other_votes

    cleaned = pd.DataFrame.from_records(records).set_index("year")
    cleaned.index = cleaned.index.astype("int64")
    cleaned.index.name = "year"
    cleaned["incumbent_party_ec_votes"] = cleaned["incumbent_party_ec_votes"].astype(
        "int64"
    )
    cleaned["challenger_party_ec_votes"] = cleaned["challenger_party_ec_votes"].astype(
        "int64"
    )
    cleaned["incumbent_is_democrat"] = cleaned["incumbent_is_democrat"].astype(bool)
    cleaned = cleaned[FINAL_COLUMNS]
    other_ec_votes = pd.Series(other_values, dtype="int64", name="other_ec_votes")
    other_ec_votes.index = other_ec_votes.index.astype("int64")
    other_ec_votes.index.name = "year"
    validate_cleaned_frame(cleaned, other_ec_votes)
    return cleaned, other_ec_votes


def validate_cleaned_frame(cleaned: pd.DataFrame, other_ec_votes: pd.Series) -> None:
    """Enforce the requested structural, type, range, and arithmetic checks."""
    if len(cleaned) != 15:
        raise AssertionError(f"Expected 15 election rows; found {len(cleaned)}.")
    if cleaned.index.tolist() != list(EXPECTED_YEARS):
        raise AssertionError("Election years are not exactly 1968-2024 in order.")
    if not cleaned.index.is_unique or not cleaned.index.is_monotonic_increasing:
        raise AssertionError("Election years must be unique and sorted ascending.")
    if cleaned.columns.tolist() != FINAL_COLUMNS:
        raise AssertionError(f"Output columns must be exactly {FINAL_COLUMNS}.")
    if cleaned.isna().any().any():
        raise AssertionError("Cleaned output contains missing values.")
    if not is_integer_dtype(cleaned["incumbent_party_ec_votes"]):
        raise AssertionError("Incumbent-party EC values are not integer typed.")
    if not is_integer_dtype(cleaned["challenger_party_ec_votes"]):
        raise AssertionError("Challenger-party EC values are not integer typed.")
    if not is_bool_dtype(cleaned["incumbent_is_democrat"]):
        raise AssertionError("incumbent_is_democrat is not boolean typed.")
    for column in FINAL_COLUMNS[:2]:
        if not cleaned[column].between(0, TOTAL_ELECTORAL_VOTES).all():
            raise AssertionError(f"{column} contains a value outside 0-538.")
    if not other_ec_votes.index.equals(cleaned.index):
        raise AssertionError("other_ec_votes is not aligned to the cleaned frame.")
    if (other_ec_votes < 0).any():
        raise AssertionError("other_ec_votes contains a negative value.")

    totals = (
        cleaned["incumbent_party_ec_votes"]
        + cleaned["challenger_party_ec_votes"]
        + other_ec_votes
    )
    if not totals.eq(TOTAL_ELECTORAL_VOTES).all():
        raise AssertionError("An election does not reconcile exactly to 538 EC votes.")
    assert (
        cleaned["incumbent_party_ec_votes"]
        + cleaned["challenger_party_ec_votes"]
        + other_ec_votes
        == TOTAL_ELECTORAL_VOTES
    ).all()


def write_and_verify(cleaned: pd.DataFrame, output_path: Path) -> None:
    """Write the requested CSV atomically and verify an exact read-back."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".csv.tmp")
    try:
        cleaned.to_csv(temporary_path, index=True, index_label="year")
        written = pd.read_csv(temporary_path)
        if written.columns.tolist() != ["year", *FINAL_COLUMNS]:
            raise AssertionError("Written CSV has unexpected columns or column order.")
        written = written.set_index("year")
        written.index = written.index.astype("int64")
        written.index.name = "year"
        written["incumbent_party_ec_votes"] = written[
            "incumbent_party_ec_votes"
        ].astype("int64")
        written["challenger_party_ec_votes"] = written[
            "challenger_party_ec_votes"
        ].astype("int64")
        if not is_bool_dtype(written["incumbent_is_democrat"]):
            raise AssertionError("CSV read-back did not preserve boolean values.")
        pd.testing.assert_frame_equal(written, cleaned, check_exact=True)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def print_validation_summary(
    cleaned: pd.DataFrame, other_ec_votes: pd.Series, output_path: Path
) -> None:
    """Print the concise per-election validation summary requested by the user."""
    print("Validation summary (Electoral College votes):")
    print("year  incumbent  challenger  other  incumbent_party")
    for year, row in cleaned.iterrows():
        party = "Democratic" if bool(row["incumbent_is_democrat"]) else "Republican"
        print(
            f"{year}  {int(row['incumbent_party_ec_votes']):9d}  "
            f"{int(row['challenger_party_ec_votes']):10d}  "
            f"{int(other_ec_votes.loc[year]):5d}  {party}"
        )
    print(f"Output: {output_path}")


def main() -> None:
    """Extract, validate, save, and summarize the presidential EC series."""
    repo_root = Path(__file__).resolve().parents[1]
    _, sources = discover_sources(repo_root)
    results = extract_all_results(sources)
    cleaned, other_ec_votes = build_cleaned_frame(results)
    output_path = repo_root / OUTPUT_RELATIVE_PATH
    write_and_verify(cleaned, output_path)
    print_validation_summary(cleaned, other_ec_votes, output_path)


if __name__ == "__main__":
    main()
